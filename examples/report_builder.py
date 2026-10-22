# -*- coding: utf-8 -*-
"""
实验报告生成器 — 支持章节、表格、代码块、截图嵌入。

用法（由 lab.ps1 或直接调用）：
    python report_builder.py --out 报告.docx --config 配置.json

配置 JSON 结构（见 build_hive_report.py 示例）：
{
  "title": "标题",
  "subtitle": "副标题",
  "meta": "日期 / 主机",
  "blocks": [
     {"type":"h1","text":"一、实验目标"},
     {"type":"h2","text":"1.1 xxx"},
     {"type":"h3","text":"Step1 xxx"},
     {"type":"p","text":"正文","bold":false},
     {"type":"bullet","text":"要点"},
     {"type":"code","text":"命令","size":9},
     {"type":"table","headers":["a","b"],"rows":[["1","2"]]},
     {"type":"image","path":"shots/yarn.png","caption":"图1 YARN 任务列表","width_cm":15.5},
     {"type":"pagebreak"}
  ]
}
"""
import argparse
import json
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"


def set_run_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run.element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def init_document(page="A4", margin_cm=2.54):
    """A4 + 统一页边距，保证 15.5cm 宽的截图放得下。"""
    doc = Document()

    for sec in doc.sections:
        if page.upper() == "A4":
            sec.page_width = Cm(21.0)
            sec.page_height = Cm(29.7)
        sec.left_margin = Cm(margin_cm)
        sec.right_margin = Cm(margin_cm)
        sec.top_margin = Cm(margin_cm)
        sec.bottom_margin = Cm(margin_cm)

    style = doc.styles["Normal"]
    style.font.name = CN_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)

    for sname in ["Heading 1", "Heading 2", "Heading 3", "Title"]:
        st = doc.styles[sname]
        st.font.name = CN_FONT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    return doc


def add_cover(doc, title, subtitle, meta):
    h = doc.add_heading(title, level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(subtitle), CN_FONT, 12, color=(0x60, 0x60, 0x60))
    if meta:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(meta), CN_FONT, 9, color=(0x80, 0x80, 0x80))


def add_code(doc, text, size=9):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    set_run_font(p.add_run(text), MONO_FONT, size, color=(0x1F, 0x3B, 0x57))
    return p


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), CN_FONT, 10.5, bold=bold)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_run_font(p.add_run(text), CN_FONT, 10.5)
    return p


def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        set_run_font(cell.paragraphs[0].add_run(str(htext)), CN_FONT, 9, bold=True)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            set_run_font(cells[i].paragraphs[0].add_run(str(val)), CN_FONT, 9)
    return t


def add_image(doc, path, caption=None, width_cm=None):
    """按页面可用宽度自动限制图片尺寸，避免溢出。"""
    if not os.path.isabs(path):
        path = os.path.abspath(path)
    if not os.path.exists(path):
        p = doc.add_paragraph()
        set_run_font(p.add_run(f"[截图缺失: {path}]"), CN_FONT, 9, color=(0xC0, 0x00, 0x00))
        return

    sec = doc.sections[0]
    avail_cm = (sec.page_width - sec.left_margin - sec.right_margin) / 360000.0
    limit = avail_cm - 0.2
    if width_cm is None or width_cm > limit:
        width_cm = limit

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(10)
        set_run_font(cp.add_run(caption), CN_FONT, 9, color=(0x60, 0x60, 0x60))


def build(cfg, out_path):
    doc = init_document()
    add_cover(doc, cfg.get("title", "实验报告"),
              cfg.get("subtitle", ""), cfg.get("meta", ""))

    base = cfg.get("base_dir") or os.getcwd()
    for blk in cfg.get("blocks", []):
        t = blk.get("type")
        if t == "h1":
            doc.add_heading(blk["text"], level=1)
        elif t == "h2":
            doc.add_heading(blk["text"], level=2)
        elif t == "h3":
            doc.add_heading(blk["text"], level=3)
        elif t == "p":
            add_para(doc, blk["text"], blk.get("bold", False))
        elif t == "bullet":
            add_bullet(doc, blk["text"])
        elif t == "code":
            add_code(doc, blk["text"], blk.get("size", 9))
        elif t == "table":
            add_table(doc, blk["headers"], blk["rows"])
        elif t == "image":
            p = blk["path"]
            if not os.path.isabs(p):
                p = os.path.join(base, p)
            add_image(doc, p, blk.get("caption"), blk.get("width_cm", 15.5))
        elif t == "pagebreak":
            doc.add_page_break()
        else:
            print(f"[warn] 未知块类型: {t}", file=sys.stderr)

    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True)
    ap.add_argument("--config", required=True)
    args = ap.parse_args()
    with open(args.config, "r", encoding="utf-8") as f:
        cfg = json.load(f)
    out = build(cfg, args.out)
    print("saved:", out)


if __name__ == "__main__":
    main()
